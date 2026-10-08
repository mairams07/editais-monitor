"""Testes do núcleo offline: planilha, normalização, deduplicação e gravação."""
import shutil
from datetime import date
from decimal import Decimal

import openpyxl
import pytest

import casamento
import montagem
import normalizacao as N
import planilha
from modelos import Cargo, Certame, Evidencia


# ------------------------------------------------------------------ normalização
@pytest.mark.parametrize("bruto,esperado", [
    ("AOCP ", "aocp"), ("Instituto AOCP", "aocp"), ("Cesgranrio ", "cesgranrio"), ("Vunesp", "vunesp"),
    ("VUNESP", "vunesp"), ("FCC", "fcc"), ("Cebraspe", "cebraspe"), ("IDECAN", "idecan"), ("IBFC", "ibfc"),
    ("Instituto ACCESS", None), ("FGV", None), ("em aberto ", None),
])
def test_banca(bruto, esperado):
    assert N.banca(bruto) == esperado


@pytest.mark.parametrize("bruto,esperado", [
    ("R$ 1.561,97", Decimal("1561.97")), ("R$ 126,79", Decimal("126.79")), ("5.197,50", Decimal("5197.50")),
    ("Gratuita", Decimal("0.00")), (220, Decimal("220.00")), ("—", None),
])
def test_dinheiro(bruto, esperado):
    assert N.dinheiro(bruto) == esperado


@pytest.mark.parametrize("orgao,esperado", [
    ("Instituto Federal de Santa Catarina", "Federal"),  # caso IFSC registrado errado como Estadual
    ("Prefeitura Municipal de Porto Velho", "Municipal"),
    ("Corpo de Bombeiros Militar de Roraima", "Estadual"),
    ("Ministério Público do Estado de São Paulo", "Estadual"),
    ("Advocacia-Geral da União", "Federal"),
])
def test_esfera(orgao, esperado):
    assert N.esfera(orgao) == esperado


# ------------------------------------------------------------------ leitura
def test_cabecalho_detectado(entrada):
    a26 = planilha.ler_aba(entrada, "2026")
    a25 = planilha.ler_aba(entrada, "2025")
    assert a26.linha_cabecalho == 2 and a25.linha_cabecalho == 1
    for c in ("CLIENTE", "BANCA", "SITUACAO", "SALARIO", "VAGAS", "TAXA", "ESFERA"):
        assert c in a26.colunas, c
    assert a26.colunas["CLIENTE"] == 5


# ------------------------------------------------------------------ deduplicação
def _cert(banca, orgao, uf, cargos=("Analista",)):
    return Certame(banca, "x", "https://exemplo", titulo="Edital 1", orgao=orgao, uf=uf,
                   publicado_em=date(2026, 3, 1), cargos=[Cargo(c) for c in cargos])


@pytest.fixture(scope="module")
def abas(entrada):
    return [planilha.ler_aba(entrada, "2025"), planilha.ler_aba(entrada, "2026")]


# Certames do gabarito: a FGV recebeu a demanda → não podem virar EXTERNO.
@pytest.mark.parametrize("banca,orgao,uf", [
    ("aocp", "Instituto Federal de Educação, Ciência e Tecnologia do Pará", "PA"),
    ("fcc", "Defensoria Pública do Estado da Paraíba", "PB"),
    ("idecan", "Corpo de Bombeiros Militar de Roraima", "RR"),
    ("idecan", "Prefeitura Municipal de Porto Velho", "RO"),
    ("fcc", "Manaus Previdência - MANAUSPREV", "AM"),
    ("cebraspe", "Escola Superior da Advocacia-Geral da União", "DF"),
    ("aocp", "Polícia Civil da Bahia", "BA"),
    ("cebraspe", "Polícia Militar de Pernambuco", "PE"),   # demanda registrada na SAD-PE (linhas 979–983)
])
def test_certame_do_gabarito_nao_vira_externo(abas, banca, orgao, uf):
    res = casamento.classificar(_cert(banca, orgao, uf), abas)
    assert res.decisao in ("JA_NA_PLANILHA", "AMBIGUO"), res


def test_policia_civil_df_nao_casa_com_policia_federal(abas):
    res = casamento.classificar(_cert("cebraspe", "Polícia Civil do Distrito Federal", "DF"), abas)
    assert all("POLICIA FEDERAL" != N.texto(c.cliente) for c in res.candidatos if c.score >= 85)


def test_orgao_generico_sem_estado_nao_casa_sozinho(abas):
    c = _cert("fcc", "Secretaria da Fazenda do Estado", "")
    assert casamento.classificar(c, abas).decisao != "JA_NA_PLANILHA"


def test_orgao_com_varios_codigos_conta_como_ja_na_planilha(abas):
    res = casamento.classificar(_cert("cebraspe", "Câmara dos Deputados", "DF"), abas)
    assert res.decisao == "JA_NA_PLANILHA"


def test_limpar_orgao():
    from extracao.campos import limpar_orgao
    assert limpar_orgao("CÂMARA MUNICIPAL DE PONTA PORÃ/MS EDITAL Nº 1 – CÂMARA", []) == "CÂMARA MUNICIPAL DE PONTA PORÃ/MS"
    assert limpar_orgao("SECRETARIA DE ESTADO DE ADMINISTRAÇÃO DO ESTADO DO", []) == "SECRETARIA DE ESTADO DE ADMINISTRAÇÃO DO ESTADO"
    assert limpar_orgao("Polícia Civil e de Escrivão de Polícia Civil da Polícia Civil do Estado de Alagoas", []) == "POLÍCIA CIVIL DO ESTADO DE ALAGOAS"
    assert limpar_orgao("Serviços Auxiliares do TCDF cumprirão jornada de trabalho", []) == ""


def test_mesmo_orgao_outra_uf_nao_casa(abas):
    res = casamento.classificar(_cert("cebraspe", "Polícia Civil do Estado do Acre", "AC"), abas)
    assert all("BAHIA" not in N.texto(c.cliente) for c in res.candidatos)


def test_orgao_inexistente_vira_externo(abas):
    res = casamento.classificar(_cert("vunesp", "Prefeitura Municipal de Xyzópolis do Norte", "SP"), abas)
    assert res.decisao == "EXTERNO"


# ------------------------------------------------------------------ montagem e gravação
def test_cliente_obrigatorio():
    with pytest.raises(montagem.CertameIncompleto):
        montagem.linhas(_cert("vunesp", "", "SP"), 2026)


def test_gravacao_em_copia(entrada, tmp_path):
    original = tmp_path / "orig.xlsx"
    shutil.copy2(entrada, original)
    hash_antes = planilha.sha256(original)
    cert = _cert("vunesp", "Prefeitura Municipal de Xyzópolis", "SP", cargos=("Agente Administrativo", "Médico"))
    cert.cargos[0].campos["SALARIO"] = Evidencia(Decimal("2500.00"), "CONFIRMADO", "Edital 1", 12, "vencimento R$ 2.500,00", "u")
    cert.campos_certame["TAXA"] = Evidencia(Decimal("60.00"), "CONFIRMADO", "Edital 1", 3, "taxa R$ 60,00", "u")
    novas = montagem.linhas(cert, 2026)
    arq, alt = planilha.gravar(original, "2026", novas, [], {}, tmp_path / "saida", date(2026, 10, 6), modo="direto")

    assert planilha.sha256(original) == hash_antes            # original intocado
    ws = openpyxl.load_workbook(arq)["2026"]
    ultima = planilha.ler_aba(original, "2026").ultima_linha_dados
    r1, r2 = ultima + 1, ultima + 2
    assert ws.cell(r1, 5).value == "PREFEITURA MUNICIPAL DE XYZÓPOLIS"  # CLIENTE preservado, caixa alta com acento
    assert ws.cell(r1, 8).value == "EXTERNO"
    assert ws.cell(r1, 3).value == "Vunesp"
    assert ws.cell(r1, 13).value == 2500.0 and ws.cell(r1, 13).fill.fgColor.rgb.endswith("C6EFCE")
    assert ws.cell(r2, 13).value is None                       # salário do 2º cargo não inventado
    assert ws.cell(r2, 17).value == 60.0                       # taxa do certame vale para todos os cargos
    assert ws.cell(r1, 2).value is None                        # CÓD_INTERNO vazio, como nas linhas EXTERNO existentes
    wb = openpyxl.load_workbook(arq)
    assert "LOG" in wb.sheetnames
    a = openpyxl.load_workbook(alt)["ALTERACOES"]
    assert a.max_row > 1 and a.cell(2, 1).value == "2026"


def test_quadro_por_cargo_cesgranrio():
    """Padrão PC-AP 2026: título do cargo + quadro 'Ampla | PcD | Cadastro de Reserva Total | Subsídio'."""
    from extracao import campos
    from extracao.pdf import Pagina
    from modelos import Certame
    p1 = Pagina(1, "EDITAL Nº 1 – PC/AP, DE 30 DE SETEMBRO DE 2026\nO Concurso Público para o cargo de Delegado de Polícia "
                   "Civil será constituído de 02 (duas) fases, a saber:\n1.2.1.1. 1ª Fase - Prova Objetiva e Prova Dissertativa, "
                   "de caráter eliminatório\n1.2.1.2. 2ª Fase - Prova Oral, de caráter classificatório\n")
    p2 = Pagina(2, "DELEGADO DE POLÍCIA CIVIL\nCadastro de\nAmpla PcD\nEscolaridade/Pré-Requisitos Reserva Subsídio\n"
                   "Concorrência (5%)\nTotal\nDiploma de graduação em Direito 97 5 102 R$ 31.439,06\n"
                   "OFICIAL INVESTIGADOR DE POLÍCIA CIVIL\nCadastro de\nAmpla PcD\nEscolaridade/Pré-Requisitos Reserva Subsídio\n"
                   "Concorrência (5%)\nTotal\nDiploma de graduação em curso superior\n279 15 294 R$ 7.327,04\n")
    p3 = Pagina(3, "no valor de R$ 200,00 (duzentos reais) para a carreira de Delegado de Polícia Civil e de R$ 150,00 "
                   "(cento e cinquenta reais) para a carreira de Oficial Investigador de Polícia Civil.\n"
                   "8.1. As Provas Objetiva e Dissertativa serão realizadas na cidade de Macapá/AP, na data prevista de 06/12/2026.\n")
    c = Certame("cesgranrio", "x", "u")
    campos.preencher_certame(c, [p1, p2, p3], "Edital", "u")
    d = {g.nome: {k: v.valor for k, v in g.campos.items()} for g in c.cargos}
    assert d["Delegado de Polícia Civil"]["VAGAS"] == "-" and d["Delegado de Polícia Civil"]["VAGAS_CR"] == 102
    assert d["Oficial Investigador de Polícia Civil"]["VAGAS_CR"] == 294
    assert str(d["Oficial Investigador de Polícia Civil"]["TAXA"]) == "150.00"
    assert d["Delegado de Polícia Civil"]["ETAPAS"] == "Prova Objetiva, Prova Discursiva, Prova Oral"
    assert c.campos_certame["CIDADES"].valor == "Macapá"


def test_quadro_tabela_aocp():
    """Quadro de vagas por linha (padrão AOCP): células com ponto inicial, cotas somadas, vencimento por classe."""
    from extracao import campos
    from extracao.pdf import Pagina
    vagas = [[".COD.", ".CARGO", ".CLASSE", ".VAGAS AC", ".PCD", ".PPP", ".TOTAL"],
             [".201", ".ASSISTENTE DE ALUNO", ".C", ".1", ".0", ".2", ".3"],
             [".402", ".ANALISTA DE TECNOLOGIA DA INFORMAÇÃO", ".E", ".1", ".0", ".1", ".2"],
             [".413", ".TÉCNICO EM ASSUNTOS ED U C A C I O N A I S", ".E", ".6", ".1", ".3", ".11"]]  # 6+1+3≠11: não preenche
    venc = [[".Nível de Classificação", ".Vencimento Básico"], [".Nível C", ".R$ 2.607,70"], [".Nível E", ".R$ 5.215,39"]]
    taxa = [["Nível", "Valor"], ["Nível E", "R$ 120,00"]]           # não é quadro de vencimento
    p = Pagina(2, "", tabelas=[vagas, venc, taxa])
    r = {x["nome"]: {k: v.valor for k, v in x.items() if hasattr(v, "valor")} for x in campos._cargos_tabela([p], "E", "u")}
    assert r["Assistente de Aluno"] == {"VAGAS": 3, "SALARIO": Decimal("2607.70")}
    assert r["Analista de Tecnologia da Informação"]["SALARIO"] == Decimal("5215.39")
    assert "VAGAS" not in r["Técnico em Assuntos Educacionais"]


def test_conferencia_llm():
    """Valor da IA só entra se o trecho existe na página e contém o valor; trecho inventado ou valor trocado caem."""
    from extracao import llm
    from extracao.pdf import Pagina
    from modelos import Certame
    pags = [Pagina(3, "DELEGADO DE POLÍCIA CIVIL\\nDiploma de graduação em Direito 97 5 102 R$ 31.439,06\\n"),
            Pagina(6, "no valor de R$ 200,00 (duzentos reais) para a carreira de Delegado de Polícia Civil e de R$ 150,00")]
    bruto = {
        "eh_edital_de_abertura": True,
        "orgao": {"valor": "Polícia Civil do Estado do Amapá", "pagina": 3, "trecho": "Polícia Civil do Estado do Amapá realizará"},
        "objeto": {"valor": None, "pagina": None, "trecho": None}, "tipo": "concurso",
        "uf": {"valor": None, "pagina": None, "trecho": None}, "cidades": {"valor": None, "pagina": None, "trecho": None},
        "etapas": {"valor": None, "pagina": None, "trecho": None}, "taxa_unica": {"valor": None, "pagina": None, "trecho": None},
        "cargos": [{
            "cargo": "Delegado de Polícia Civil", "especialidade": None,
            "nivel": {"valor": "Superior", "pagina": 3, "trecho": "Diploma de graduação em Direito"},
            "salario": {"valor": "31439.06", "pagina": 3, "trecho": "97 5 102 R$ 31.439,06"},
            "vagas": {"valor": "-", "pagina": 3, "trecho": "97 5 102"},
            "vagas_cr": {"valor": "102", "pagina": 3, "trecho": "97 5 102 R$ 31.439,06"},
            "taxa": {"valor": "250.00", "pagina": 6, "trecho": "no valor de R$ 200,00 (duzentos reais)"},
            "etapas": {"valor": None, "pagina": None, "trecho": None}}],
    }
    c = Certame("cesgranrio", "x", "u", orgao="")
    descartes = llm.aplicar(c, bruto, pags, "Edital", "u")
    cg = c.cargos[0].campos
    assert c.orgao == ""                                   # trecho inventado: não existe na página
    assert str(cg["SALARIO"].valor) == "31439.06" and cg["VAGAS_CR"].valor == 102 and cg["NIVEL"].valor == "Superior"
    assert "TAXA" not in cg                                # 250,00 não está no trecho
    motivos = {d["campo"]: d["motivo"] for d in descartes}
    assert motivos["CLIENTE"] == "trecho não confere com a página" and motivos["TAXA"] == "valor não aparece no trecho"


def test_gravacao_modo_revisao(entrada, tmp_path):
    """Modo revisão: aba do ano intocada; linhas em A_CONFERIR com o mesmo cabeçalho e as mesmas posições de coluna."""
    import openpyxl
    original = tmp_path / "orig.xlsx"
    shutil.copy2(entrada, original)
    cert = _cert("vunesp", "Prefeitura Municipal de Xyzópolis", "SP", cargos=("Analista",))
    cert.cargos[0].campos["SALARIO"] = Evidencia(Decimal("2500.00"), "CONFIRMADO", "Edital 1", 12, "vencimento R$ 2.500,00", "u")
    novas = montagem.linhas(cert, 2026)
    aba0 = planilha.ler_aba(original, "2026")
    arq, _ = planilha.gravar(original, "2026", novas, [], {}, tmp_path / "saida", date(2026, 10, 8))
    assert planilha.ler_aba(arq, "2026").ultima_linha_dados == aba0.ultima_linha_dados   # aba do ano sem linha nova
    wb = openpyxl.load_workbook(arq)
    cab_ano = [c.value for c in wb["2026"][aba0.linha_cabecalho]]
    ac = wb["A_CONFERIR"]
    assert [c.value for c in ac[1]][:len(cab_ano)] == cab_ano
    assert ac.cell(2, aba0.colunas["CLIENTE"]).value == "PREFEITURA MUNICIPAL DE XYZÓPOLIS"
    assert ac.cell(2, aba0.colunas["SALARIO"]).value == 2500.0
    assert "vencimento R$ 2.500,00" in (ac.cell(2, max(aba0.colunas.values()) + 4).value or "")
